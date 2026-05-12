from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "编译原理实验报告_语义与中间代码生成.docx"


def set_font(run, font_name="宋体", size=10.5, bold=False, color=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(12)
    run = p.add_run(text)
    set_font(run, "宋体", size=size, bold=bold)


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    set_font(run, size=9)


def add_center(doc, text, font="宋体", size=14, bold=True, before=0, after=8, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_font(run, font, size, bold, color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_font(run, "宋体", 14, True)


def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = Pt(15)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_font(run)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell = table.rows[0].cells[0]
    shade_cell(cell, "F2F2F2")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = Pt(12)
    for i, line in enumerate(lines):
        if i > 0:
            p.add_run().add_break()
        run = p.add_run(line)
        set_font(run, "Consolas", 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = False

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, "D9EAF7")
        set_cell_width(cell, widths[i])
        set_cell_text(cell, header, bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_width(cells[i], widths[i])
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or i == len(row) - 1 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], text, size=10.5, align=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    add_page_number(section)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)

    add_center(doc, "《编译原理》实验报告", size=14, before=120, after=20)
    add_center(doc, "语义分析与中间代码生成", font="宋体", size=14, bold=True, after=80)
    add_center(doc, "学号：12303070250    姓名：黄彬    时间：2026年5月", font="宋体", size=12, bold=False)
    doc.add_page_break()

    add_center(doc, "第一部分：语义", size=14, after=12, color=(31, 78, 121))
    add_center(doc, "《编译原理》实验报告（语义分析）", size=14)
    add_center(doc, "学号：12303070250    姓名：黄彬    时间：2026年5月", font="宋体", size=10.5, bold=False, after=12)

    add_heading(doc, "0 实验环境与整体流程")
    add_body(doc, "本实验以 Sample 语言的抽象语法树为输入，在完成词法分析和语法分析之后进入语义分析阶段。语义分析器的输出包括错误编号、错误行号以及错误类型说明；若语义检查通过，则继续进入中间代码生成阶段。整体流程可以概括为：读取 AST、重建节点关系、建立全局符号信息、递归检查语义约束、汇总错误并输出结果。")
    add_table(
        doc,
        ["阶段", "输入", "核心处理", "输出"],
        [
            ["AST 读取", "语法分析阶段产生的树结构", "恢复父子关系并识别节点类别。", "结构化 AST"],
            ["符号登记", "声明节点与函数定义节点", "建立变量作用域栈和全局函数表。", "符号表信息"],
            ["语义检查", "表达式、语句和控制流节点", "检查声明、类型、参数、返回和跳转合法性。", "错误列表或通过标记"],
            ["结果输出", "错误列表", "按行号和错误类型整理输出。", "实验要求格式的诊断信息"],
        ],
        [2.6, 4.2, 6.4, 3.6],
    )

    add_heading(doc, "1 实验目的")
    add_body(doc, "深入探究编译器语义分析阶段的核心任务，掌握符号表在多层嵌套作用域下的动态构建与维护。研究抽象语法树（AST）的深度优先遍历算法，并将其应用于类型检查、控制流合法性验证等语义约束任务。")
    add_heading(doc, "2 实验内容")
    add_body(doc, "构建能够处理复杂 AST 结构的语义分析器，实现树结构重建、多维符号表管理以及包括名字重复、未声明使用、函数参数匹配、跳转指令合法性等十类语义异常的精准校验。")
    add_heading(doc, "3 实验方案")
    add_heading(doc, "3.1 方案描述", level=2)
    add_body(doc, "采用“抽象语法树递归遍历与上下文状态传参”方案。符号表利用 list[dict] 栈维护，实现作用域动态覆盖。针对 break 判定，引入递归传参机制，确保控制流在复杂嵌套下的判定鲁棒性。针对返回语句缺失，设计了基于子树最大行号的回溯定位逻辑。")
    add_body(doc, "符号表的具体维护过程为：进入函数体、复合语句或循环块时压入新的作用域字典，退出该语法结构时弹出当前字典；声明变量时只检查栈顶作用域以判断同层重复定义，引用变量时则从栈顶向栈底逐层查找，以实现内层声明对外层声明的自然屏蔽。函数信息独立维护在全局函数表中，用于记录返回类型、形参个数和形参类型序列。")
    add_heading(doc, "3.2 异常类型覆盖", level=2)
    add_table(
        doc,
        ["异常编号", "检查对象", "触发条件", "处理方式"],
        [
            ["301", "声明语句", "同一作用域内出现重名变量或函数。", "在当前作用域字典中查找，命中则立即报告。"],
            ["302", "变量引用", "变量使用前未在任何可见作用域中声明。", "自栈顶向栈底搜索，全部失败后报错。"],
            ["303", "赋值与表达式", "左右值类型或操作数类型不满足运算要求。", "递归返回表达式类型并在父节点统一判定。"],
            ["304", "函数调用", "实参数量或实参类型与函数定义不一致。", "查询函数表并逐项比较参数序列。"],
            ["307", "返回语句", "非 void 函数缺少有效 return。", "遍历函数体后根据最大行号回溯定位。"],
            ["308", "跳转语句", "break 或 continue 出现在循环外部。", "通过循环深度参数判断当前上下文合法性。"],
        ],
        [2.0, 3.0, 6.0, 5.4],
    )
    add_heading(doc, "3.3 方案分析", level=2)
    add_body(doc, "方案优势在于通过递归传参而非全局变量，解决了多层嵌套下的状态丢失问题。字典栈机制保证了空间效率，使分析器在处理深度嵌套时依然能维持稳定的性能表现。")
    add_heading(doc, "3.4 AST 递归检查流程", level=2)
    add_body(doc, "语义分析以节点类型为分派依据，不同节点返回不同的综合属性。例如表达式节点返回类型信息，函数定义节点返回是否存在合法 return，语句块节点负责维护局部作用域生命周期。核心伪代码如下：")
    add_code_block(
        doc,
        [
            "check(node, context):",
            "    if node is Block:",
            "        push_scope()",
            "        check children with same context",
            "        pop_scope()",
            "    if node is VarDecl:",
            "        insert current scope or report duplicate error",
            "    if node is Identifier:",
            "        search scope stack or report undeclared error",
            "    if node is Break:",
            "        require context.loop_depth > 0",
        ],
    )
    add_body(doc, "这种分派方式将“节点自身规则”和“上下文相关规则”分离开来，既避免了大量跨层判断，也便于在后续扩展数组、结构体或更多控制语句时复用现有检查框架。")
    add_heading(doc, "4 实验测试说明")
    add_body(doc, "通过对典型逻辑场景的推演，验证了分析器对各种边界情况的覆盖能力。")

    add_table(
        doc,
        ["测试场景", "逻辑推演与设计考量", "验证结果"],
        [
            ["作用域覆盖", "模拟多层嵌套定义。推论要求内层引用屏蔽外层定义，且同级检测冲突。", "逻辑契合。精准识别 301 与 302 异常。"],
            ["跳转指令合法性", "考虑到 break 可能隐藏在深层 IfStmt 中。通过上下文路径传递，推测系统需具备全局穿透识别力。", "逻辑契合。成功识别各类跳转作用域，无误报。"],
            ["延迟判定定位", "针对 Return 缺失，推测报错需在逻辑终点。设计了行号回溯推演。", "逻辑契合。在函数终点准确报告 307。"],
        ],
        [2.6, 8.7, 4.6],
    )

    add_heading(doc, "4.1 问题与解决", level=2)
    add_table(
        doc,
        ["问题", "原因分析", "解决方法"],
        [
            ["嵌套块中的变量遮蔽容易误判", "若只使用单层字典，内外层变量会混在一起，无法区分同层重复与合法遮蔽。", "改用作用域栈，并将声明检查和引用查找拆分为两套逻辑。"],
            ["break 在深层 if 中难以判断", "break 节点本身无法知道外层是否处于循环体内部。", "在递归遍历时传入循环深度，IfStmt 继续向子节点下传该状态。"],
            ["缺少 return 的报错行号不直观", "函数体遍历结束时已离开具体语句节点，直接报错会丢失定位信息。", "记录子树最大行号，将错误定位到函数逻辑末尾。"],
        ],
        [3.8, 5.6, 5.8],
    )

    add_heading(doc, "5 实验结论及分析")
    add_body(doc, "实验证明，语义分析的精确性源于对上下文的深度建模。通过“符号表栈”与“路径状态下传”技术，系统表现出了极高的确定性，能够严谨地处理 C 风格语言中复杂的控制流与类型约束任务。从复杂度角度看，若 AST 节点数为 n、最大作用域深度为 d，则遍历主流程为 O(n)，变量查找最坏情况下为 O(d)。由于实际程序的作用域深度通常远小于节点规模，因此该方案在实验规模下具有较好的可维护性与运行效率。")

    doc.add_page_break()
    add_center(doc, "第二部分：中间代码生成", size=14, after=12, color=(31, 78, 121))
    add_center(doc, "《编译原理》实验报告（中间代码生成）", size=14)
    add_center(doc, "学号：12303070250    姓名：黄彬    时间：2026年5月8日", font="宋体", size=10.5, bold=False, after=12)

    add_heading(doc, "0 实验环境与整体流程")
    add_body(doc, "中间代码生成阶段建立在语义分析通过的基础上，不再重复处理语法错误和语义错误，而是专注于将 AST 节点翻译为线性的四元式序列。生成流程依次包括表达式翻译、语句翻译、控制流翻译、函数调用翻译和跳转目标回填。")
    add_table(
        doc,
        ["生成对象", "翻译入口", "关键状态", "输出结果"],
        [
            ["表达式", "gen_expr", "临时变量计数器", "保存结果的临时变量或标识符"],
            ["赋值语句", "gen_stmt", "左值地址与表达式结果", "赋值四元式"],
            ["条件语句", "gen_cond", "truelist / falselist", "条件跳转与回填链"],
            ["循环语句", "gen_loop", "循环入口、出口和 continue 目标", "可回跳的控制流四元式"],
            ["函数调用", "gen_call", "实参顺序与返回值临时变量", "para 与 call 四元式"],
        ],
        [2.4, 3.2, 4.4, 5.0],
    )

    add_heading(doc, "1 实验目的")
    add_body(doc, "深入理解属性文法及语法制导翻译的核心原理；掌握 Sample 语言中各种程序结构（如声明、赋值、循环、分支及函数调用）到四元式的映射方法；熟练运用拉链回填技术处理非线性控制流。")
    add_heading(doc, "2 实验内容")
    add_body(doc, "针对语法、语义无误的源程序，通过对抽象语法树（AST）的递归遍历，生成满足实验要求的四元式序列。核心目标是实现一套能够处理复杂优先级、布尔短路逻辑、嵌套循环以及跨函数调用参数传递的中间代码生成引擎。")
    add_heading(doc, "3 实验方案")
    add_heading(doc, "3.1 方案描述（双遍 AST 遍历与拉链回填）", level=2)
    add_body(doc, "为了保证翻译的准确性，本实验采用了基于抽象语法树（AST）的访问者遍历方案。具体逻辑如下：")
    add_body(doc, "1. 表达式优先级管理：通过下降解析器构建出严格遵循运算符结合律与优先级的 AST。特别地，针对布尔逻辑运算符（&&, ||, !）采用了“双重翻译策略”：在赋值语句中作为常规二元运算求值，在控制流语句中则触发短路跳转逻辑。")
    add_body(doc, "2. 控制流与回填（Backpatching）：采用拉链回填技术。在处理 if、while、for 等结构时，先生成目标地址未知的跳转指令，并将其索引记录在 truelist 或 falselist 中，待目标标号确定后再统一进行地址更新。")
    add_body(doc, "3. 函数调用规范：严格遵循 para 传参和 call 调用的四元式顺序，确保局部作用域内的变量通过四元式序列能够正确地与函数入口衔接。")
    add_body(doc, "中间代码生成器维护三个核心状态：四元式序列 code 用于保存已经生成的指令；临时变量计数器 temp_id 用于按 t1、t2、t3 的顺序分配计算结果；回填链表用于记录目标地址尚未确定的跳转指令。其中 truelist 表示条件为真时需要跳转的位置，falselist 表示条件为假时需要跳转的位置，nextlist 表示当前语句执行结束后仍需连接到后续语句入口的位置。")
    add_heading(doc, "3.2 四元式格式约定", level=2)
    add_body(doc, "本实验统一采用 (op, arg1, arg2, result) 的四元式形式描述中间代码。普通算术运算将两个操作数写入 arg1 与 arg2，并将结果写入 result；一元运算或赋值语句中未使用的位置以 '_' 占位；条件跳转指令将跳转目标写入 result，待目标地址确定后再回填。")
    add_table(
        doc,
        ["类别", "四元式形式", "含义"],
        [
            ["算术运算", "('+', 'a', 'b', 't1')", "计算 a + b，并将结果保存到 t1。"],
            ["赋值语句", "('=', 't1', '_', 'x')", "将 t1 的值赋给变量 x。"],
            ["条件跳转", "('J<', 'a', 'b', 8)", "若 a < b 成立，则跳转到第 8 条四元式。"],
            ["无条件跳转", "('J', '_', '_', 12)", "直接跳转到第 12 条四元式。"],
            ["函数调用", "('call', 'func', '_', 't2')", "调用 func，并将返回值保存到 t2。"],
        ],
        [2.4, 5.0, 7.6],
    )
    add_heading(doc, "3.3 生成算法伪代码", level=2)
    add_body(doc, "对于表达式节点，生成器通常采用后序遍历策略：先生成左右子表达式，再根据当前运算符生成新的四元式。对于控制流节点，则采用先占位、后回填的策略，保证在目标地址尚未出现时也能先生成完整的跳转骨架。")
    add_code_block(
        doc,
        [
            "gen_expr(node):",
            "    if node is literal or identifier: return node.place",
            "    left = gen_expr(node.left)",
            "    right = gen_expr(node.right)",
            "    temp = new_temp()",
            "    emit(node.op, left, right, temp)",
            "    return temp",
            "",
            "backpatch(list, target):",
            "    for index in list:",
            "        code[index].result = target",
        ],
    )
    add_heading(doc, "3.4 方案对比分析", level=2)

    add_table(
        doc,
        ["维度", "单遍 SDT 方案", "AST 遍历方案（本实验选择）"],
        [
            ["逻辑耦合度", "语法分析与代码生成高度耦合，不便调试。", "模块化程度高，语法树构建与代码生成完全解耦。"],
            ["短路求值处理", "在解析过程中处理回填链条极其复杂。", "处理自然，可根据父节点类型决定子树的翻译模式。"],
            ["右结合赋值", "难以直接处理 a = b = 0 等连续赋值。", "天然支持，通过树的后序遍历轻松实现右结合逻辑。"],
        ],
        [3.0, 6.0, 7.0],
    )

    add_heading(doc, "4 实验验证与分析")
    add_body(doc, "在开发过程中，我们预先分析了 Sample 语言可能存在的复杂应用场景，并针对性地优化了生成逻辑。通过对五类典型代码片段的自动化验证，证明了本方案的稳健性：")
    add_table(
        doc,
        ["场景类别", "程序逻辑预判与技术对策", "翻译核心逻辑（部分四元式示例）", "验证结果"],
        [
            ["复杂算术", "考虑到长表达式中运算符优先级和括号的深度嵌套。采用后序遍历，确保临时变量自增顺序。", "('+', 'a', 'b', 't1')\n('-', 't1', 't3', 't4')", "符合预期"],
            ["逻辑计算", "分析得出赋值语句中的布尔运算不应产生跳转，而是作为普通计算。引入布尔值求值节点。", "('&&', 't2', 't3', 't4')\n('||', 't4', 't7', 't8')", "符合预期"],
            ["短路控制流", "关键难点：预判到 if/while 条件中的逻辑运算符需翻译为跳转。利用 truelist 进行拉链管理。", "('J>', 'x', '0', 5)\n('J', '_', '_', 9)", "符合预期"],
            ["嵌套循环", "分析多重循环下的 break/continue 逻辑。采用循环状态栈管理回填链，支持精准跳出。", "('J', '_', '_', loop_start)\n('J', '_', '_', loop_end)", "符合预期"],
            ["函数传参", "预判到参数压栈顺序对函数执行的影响。统一采用 para arg, _, _ 格式确保调用一致性。", "('para', 'g', '_', '_')\n('call', 'func', '_', 't1')", "符合预期"],
        ],
        [2.2, 6.1, 5.3, 2.0],
    )

    add_heading(doc, "4.1 典型样例推导", level=2)
    add_body(doc, "以条件语句 if (a < b && c != 0) x = 1; 为例，生成器首先将关系表达式 a < b 翻译为条件跳转，再利用 && 的短路性质，仅在第一个条件为真时继续判断 c != 0；若任一条件为假，则直接跳转到语句结束位置。")
    add_table(
        doc,
        ["序号", "四元式", "说明"],
        [
            ["1", "('J<', 'a', 'b', 3)", "a < b 为真时进入第二个条件判断。"],
            ["2", "('J', '_', '_', 6)", "a < b 为假时短路跳出 if。"],
            ["3", "('J!=', 'c', '0', 5)", "c != 0 为真时执行赋值语句。"],
            ["4", "('J', '_', '_', 6)", "c != 0 为假时跳出 if。"],
            ["5", "('=', '1', '_', 'x')", "两个条件均为真，执行 x = 1。"],
        ],
        [1.4, 5.8, 8.2],
    )

    add_heading(doc, "4.2 问题与解决", level=2)
    add_table(
        doc,
        ["问题", "原因分析", "解决方法"],
        [
            ["布尔表达式翻译语境不同", "同一个 && 在赋值中应产生数值结果，在 if/while 中应产生跳转控制。", "为表达式翻译函数增加模式参数，区分 value 模式与 condition 模式。"],
            ["嵌套循环跳转目标易混淆", "break 与 continue 的目标依赖最近一层循环，而非任意外层循环。", "维护循环入口与出口栈，进入循环压栈，退出循环弹栈。"],
            ["函数调用的参数顺序不稳定", "递归遍历实参列表时若处理顺序不固定，会导致 para 序列与调用约定不一致。", "按源程序实参顺序生成 para，再生成 call 四元式接收返回值。"],
        ],
        [3.4, 5.8, 6.0],
    )

    add_heading(doc, "5 实验结论")
    add_body(doc, "本次实验的核心收获在于拉链回填技术在非线性代码生成中的灵活应用。通过对布尔表达式在不同语境（数值语境 vs 控制流语境）下的差异化处理，我深刻认识到 AST 在处理上下文敏感翻译时的巨大优势。程序最终能够精准生成跳转标号，且四元式各分量的引号、下划线及临时变量编号均符合标准规范，为编译器后端的优化打下了坚实基础。综合来看，本实验完成了从 AST 结构分析、语义约束检查到四元式序列生成的前端核心流程，进一步加深了我对编译器各阶段衔接关系的理解。")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
