from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = "毕昇编译器SSA形式IR拓展阅读.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 8 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(11)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "黑体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    r.font.size = Pt(15 if level == 1 else 13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(31, 78, 121)
    return p


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F6FA")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(lines):
        if i:
            p.add_run("\n")
        r = p.add_run(line)
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        r.font.size = Pt(9.5)
    doc.add_paragraph()


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.4)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.6)
section.right_margin = Cm(2.6)

styles = doc.styles
styles["Normal"].font.name = "宋体"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Normal"].font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(70)
title.paragraph_format.space_after = Pt(24)
run = title.add_run("拓展阅读：毕昇编译器的中间表示")
run.font.name = "黑体"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
run.font.size = Pt(22)
run.font.bold = True
run.font.color.rgb = RGBColor(31, 78, 121)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(36)
r = subtitle.add_run("SSA 形式 IR 的特点及其优化优势")
r.font.name = "宋体"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
r.font.size = Pt(14)

info_table = doc.add_table(rows=4, cols=2)
info_table.style = "Table Grid"
for row in info_table.rows:
    row.height = Cm(0.9)
labels = ["课程名称", "学生姓名", "班级/学号", "提交日期"]
values = ["编译原理", "（请填写）", "（请填写）", "2026 年 5 月"]
for i, (label, value) in enumerate(zip(labels, values)):
    set_cell_text(info_table.cell(i, 0), label, bold=True)
    set_cell_shading(info_table.cell(i, 0), "D9EAF7")
    set_cell_text(info_table.cell(i, 1), value)

doc.add_section(WD_SECTION.NEW_PAGE)

add_heading(doc, "一、引言")
add_body_paragraph(
    doc,
    "毕昇编译器是面向鲲鹏生态的重要编译工具链。根据鲲鹏社区文档，毕昇编译器基于 LLVM 开发，并集成了 C/C++/Fortran 等语言前端及编译驱动；这意味着它的中间表示与优化体系继承了 LLVM IR 的核心思想。LLVM 官方语言参考也明确说明，LLVM IR 采用静态单赋值形式（SSA form）。因此，理解 SSA 形式的 IR，是理解毕昇编译器数据流分析和优化机制的一个关键入口。"
)

add_heading(doc, "二、SSA 形式 IR 的基本思想")
add_body_paragraph(
    doc,
    "SSA 的全称是 Static Single Assignment，即静态单赋值。它要求程序中的每一个变量名只被赋值一次；如果源程序中同一个变量在不同位置被多次更新，编译器会在 IR 中把它们重命名为不同的版本。这样，变量名本身就带有“定义点”的信息，编译器可以直接从一次定义追踪到所有使用点。"
)
add_body_paragraph(
    doc,
    "控制流汇合处需要特殊处理。例如一个变量可能来自 if 分支的两个不同赋值，SSA 会使用 phi 节点表示“根据前驱基本块选择哪个值”。phi 节点不是源程序中的普通语句，而是 SSA IR 在控制流图上表达数据来源的一种机制。"
)

add_code_block(
    doc,
    [
        "传统写法：",
        "x = 1",
        "if (cond) x = 2",
        "y = x + 3",
        "",
        "SSA 思路：",
        "x1 = 1",
        "x2 = 2",
        "x3 = phi(x1, x2)",
        "y1 = x3 + 3",
    ],
)

add_heading(doc, "三、与传统三地址码的差异")
add_body_paragraph(
    doc,
    "传统三地址码通常把程序拆成形如“a = b op c”的语句序列，适合表达基本运算和临时变量，但同一个变量可能在不同位置反复被定义。编译器若要判断某次使用到底来自哪次定义，就必须额外进行到达定义、活跃变量、使用-定义链等分析。SSA 则把这些关系显式化：一个名字只对应一个定义，使用点可以自然回溯到定义点，控制流合并则通过 phi 节点集中表达。"
)

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
headers = ["比较维度", "传统三地址码", "SSA 形式 IR"]
for j, header in enumerate(headers):
    set_cell_text(table.cell(0, j), header, bold=True)
    set_cell_shading(table.cell(0, j), "D9EAF7")
rows = [
    ("变量定义", "同一变量可多次赋值", "每个变量版本只定义一次"),
    ("数据依赖", "需要额外构建 def-use/use-def 信息", "定义与使用关系更直接"),
    ("控制流汇合", "通常隐含在变量值变化中", "用 phi 节点显式表达"),
    ("优化便利性", "分析前置工作较多", "更利于稀疏数据流分析和局部化变换"),
]
for row in rows:
    cells = table.add_row().cells
    for j, value in enumerate(row):
        set_cell_text(cells[j], value, bold=(j == 0))

add_heading(doc, "四、SSA 在数据流分析中的优势")
add_body_paragraph(
    doc,
    "第一，SSA 简化了定义-使用关系。由于每个 SSA 名字只有一个定义点，编译器不必在大量可能的赋值语句之间反复求解“这个值从哪里来”。这会让常量传播、复制传播、死代码删除等分析更直接，也降低了数据流方程求解的复杂度。"
)
add_body_paragraph(
    doc,
    "第二，SSA 更适合稀疏数据流分析。传统数据流分析常以基本块或程序点为单位传播集合，可能会扫描很多与目标变量无关的位置。SSA 中的 def-use 链把分析范围约束在真正相关的值和使用点上，因此分析可以沿着值流动的边进行，避免无效传播。"
)
add_body_paragraph(
    doc,
    "第三，SSA 让控制流和数据流的交汇关系更清楚。phi 节点把不同分支进入同一基本块时的值选择显式化，使编译器能够更准确地判断变量在循环、分支和合并路径中的来源。这一点对循环优化、条件分支优化和全局值编号尤其重要。"
)

add_heading(doc, "五、SSA 对优化的帮助")
add_body_paragraph(
    doc,
    "在常量传播中，如果某个 SSA 值被证明为常量，那么它的所有使用点都可以直接替换；在死代码删除中，如果一个 SSA 定义没有有效使用，且该定义没有副作用，就可以安全删除；在公共子表达式消除和全局值编号中，SSA 的单定义特性也让表达式等价性判断更清晰。"
)
add_body_paragraph(
    doc,
    "对于现代优化器来说，SSA 不只是语法形式，而是一种把程序变成“值流图”的组织方式。毕昇编译器基于 LLVM 技术路线，因此它能够利用 LLVM IR 与优化 Pass 体系，在目标平台相关优化之前进行大量平台无关优化。SSA 形式在这里起到基础支撑作用：它提高了分析精度，也让优化结果更容易保持语义正确。"
)

add_heading(doc, "六、小结")
add_body_paragraph(
    doc,
    "总的来说，毕昇编译器采用的 SSA 形式 IR 相比传统三地址码，更强调“值的来源”和“值的使用”之间的明确关系。三地址码适合作为线性化的中间代码表达，而 SSA 更适合作为优化器内部的分析与变换基础。通过变量版本化和 phi 节点，SSA 将数据依赖显式化，使数据流分析更稀疏、更精确，也使常量传播、死代码删除、公共子表达式消除、循环优化等编译优化更容易实现和组合。"
)

add_heading(doc, "参考资料")
refs = [
    "鲲鹏社区文档：《毕昇编译器介绍》，https://www.hikunpeng.com/document/detail/zh/kunpengdevkithistory/bisheng/hist-bisheng/kunpengbisheng_06_0001_1.html",
    "LLVM Project: LLVM Language Reference Manual，https://llvm.org/docs/LangRef.html",
    "华保健、高耀清：《毕昇编译器原理与实践》，清华大学出版社，2022 年。",
]
for item in refs:
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-18)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(item)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)

doc.save(OUTPUT)
print(OUTPUT)
