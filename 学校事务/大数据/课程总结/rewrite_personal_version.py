from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path("12303070250黄彬.docx")


def style_run(run, size=12, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_para(doc, text="", size=12, bold=False, align=None, first=False, before=0, after=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if first:
        p.paragraph_format.first_line_indent = Pt(24)
    run = p.add_run(text)
    style_run(run, size=size, bold=bold)
    return p


def set_cell(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    style_run(run, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def borderless(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement("w:" + edge)
        elem.set(qn("w:val"), "nil")
        borders.append(elem)
    table._tbl.tblPr.append(borders)


sections = [
    ("一、对大数据技术的理解", [
        "这门课刚开始的时候，我对大数据的理解其实挺模糊。听到“大数据”三个字，第一反应就是数据特别多，可能电脑配置要好一点，程序跑得久一点。后来学到数据采集、存储、清洗和计算这些内容，我才发现它不是简单把很多数据放在一起。真正麻烦的地方，是数据又多又乱，而且很多时候并不是拿来就能用。",
        "比如一份表里，字段名可能不统一，有的地方空着，有的记录重复，还有一些值看起来就不太合理。以前我看到这种情况，会想着先把代码写出来再说。但现在回头看，如果前面没处理好，后面算出来的结果很可能是偏的。课堂上讲到数据预处理时，我一开始觉得这一块有点琐碎，后来做练习才发现，它反而是最容易出问题的地方。",
        "我对分布式处理的印象也比较深。之前写程序，大多数时候都是默认一台电脑完成全部任务。学到 HDFS 和 MapReduce 后，我才知道大数据处理更强调把任务拆开。数据分块存储，计算也分成很多小任务，最后再汇总。这个思路不只是技术上的安排，对我理解复杂问题也有帮助。一个问题太大时，先拆开，比一上来硬做要清楚很多。",
        "不过我现在对这些技术也只是入门理解。像 Hadoop 生态里还有很多组件，我能说出大概作用，但还不能很熟练地搭建和调优。这个地方也是我后面需要补的。至少通过这门课，我知道了大数据处理不是一个单独的软件，而是一整套流程和思维方式。",
    ]),
    ("二、对应用价值和就业前景的理解", [
        "学这门课之前，我觉得大数据应用离自己有点远，好像是互联网公司或者大型企业才会用。后来发现，其实平时生活里到处都有。购物平台推荐商品，地图软件判断拥堵，短视频平台推荐内容，银行判断异常交易，这些背后都离不开数据。只是平时用的时候不会特别去想。",
        "我觉得大数据最大的价值，是让很多判断有了依据。以前可能靠经验说“这个商品可能受欢迎”，现在可以看浏览量、转化率、用户停留时间等数据。当然，数据也不能代替人做所有判断。比如推荐系统如果只盯着点击率，可能会让用户一直看到同一类内容；如果金融风控只看模型结果，也可能误伤正常用户。所以我觉得大数据应用不仅要追求效率，还要考虑场景和后果。",
        "从就业角度看，大数据相关岗位确实不少，但我也意识到它不是学完一门课就能胜任的。数据开发需要会数据库、编程和平台工具；数据分析要能看懂业务，不能只会画图；算法岗位要求更高，还要有数学和模型基础。对我来说，这门课更像是打开了一个入口，让我知道以后如果想往这个方向走，应该继续补哪些能力。",
        "还有一点我以前没太注意，就是表达能力也很重要。数据处理完以后，如果不能把结果讲清楚，别人也很难相信或者使用这个结果。以后不管是做作业还是做项目，我都应该多想一步：这个结果说明了什么，有没有可能解释错，能不能让别人看懂。",
    ]),
    ("三、个人课程收获、思考和困惑", [
        "这学期最大的收获，是我处理数据时比以前有顺序了。以前拿到数据就急着写代码，觉得先跑起来再说。现在我会先看字段、类型、缺失值和重复记录。这个习惯虽然很普通，但能少走很多弯路。有时候问题不在代码，而在数据本身。",
        "我记得做练习时，有一次只是想统计数量，结果算出来明显偏大。后来检查才发现有重复记录没有处理。那次以后我对“程序能运行”和“结果可信”这两件事分得更清楚了。程序没有报错，只能说明语法和运行环境暂时没问题，并不能说明分析就对。",
        "还有一些基础问题也暴露出来。比如路径写错、文件名没对应、编码不一致，这些小问题都可能让程序跑不起来。刚开始遇到报错我会比较急，后来慢慢学会先看报错信息，再检查文件位置、字段名和数据格式。虽然这个过程有点慢，但比乱改代码有效。",
        "我也开始意识到，数据分析里面有很多主观选择。比如什么算活跃用户，用登录次数算，还是用使用时长算，结论可能不一样。指标怎么定，异常值删不删，缺失值怎么补，这些都不是机械操作。做数据的人如果不认真，结果看起来很客观，其实可能已经带了偏差。",
        "这也让我反思自己以前写总结或者做实验报告时，常常只写最后结果，很少写中间判断。现在再看，过程其实也很重要。为什么这样处理数据，为什么选择这个指标，哪里可能不准确，这些内容写出来，反而能说明自己真的理解了。",
        "目前我还有一些困惑。真实企业里的数据肯定比课堂数据复杂得多，数据安全和隐私怎么保证？模型算出来的结果如果和人的经验不一样，应该怎么判断？还有那么多大数据框架，实际项目里到底怎么选？这些问题我现在还没有答案，只能说有了一个方向。",
        "总的来说，这门课让我对大数据有了比较实际的认识。它不是只讲概念，也不是简单运行工具，而是从数据到结果的一整套过程。以后即使我不一定直接做大数据岗位，这种先整理信息、再分析问题、最后谨慎解释结果的思路，对学习和工作应该都有帮助。",
    ]),
]


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.54)
sec.bottom_margin = Cm(2.54)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(2.6)
doc.styles["Normal"].font.name = "宋体"
doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
doc.styles["Normal"].font.size = Pt(12)

for _ in range(5):
    add_para(doc)
add_para(doc, "《大数据处理与实践》", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "课程总结", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(5):
    add_para(doc)

table = doc.add_table(rows=6, cols=2)
table.alignment = WD_ALIGN_PARAGRAPH.CENTER
borderless(table)
for row, (label, value) in zip(table.rows, [
    ("班级", "【请填写班级】"),
    ("学号", "12303070250"),
    ("姓名", "黄彬"),
    ("教师", "成 卫"),
    ("", ""),
    ("日期", "2025年  月   日"),
]):
    row.height = Cm(0.95)
    set_cell(row.cells[0], label, bold=bool(label))
    set_cell(row.cells[1], value)
    row.cells[0].width = Cm(3.0)
    row.cells[1].width = Cm(8.5)

doc.add_page_break()
add_para(doc, "课程总结和心得", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
for title, paragraphs in sections:
    add_para(doc, title, bold=True, before=8, after=2)
    for text in paragraphs:
        add_para(doc, text, first=True)

doc.save(OUT)
plain = "\n".join(p for _, paras in sections for p in paras)
print(f"saved={OUT.resolve()}")
print(f"cn_chars={len(re.findall(r'[\u4e00-\u9fff]', plain))}")
