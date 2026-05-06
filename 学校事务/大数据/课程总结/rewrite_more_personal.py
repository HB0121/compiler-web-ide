from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path("12303070250黄彬.docx")


def style_run(run, size=12, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def p(doc, text="", size=12, bold=False, align=None, first=False, before=0, after=0):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    if first:
        para.paragraph_format.first_line_indent = Pt(24)
    run = para.add_run(text)
    style_run(run, size=size, bold=bold)
    return para


def set_cell(cell, text, bold=False):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.line_spacing = 1.5
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text)
    style_run(run, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def no_border(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement("w:" + edge)
        elem.set(qn("w:val"), "nil")
        borders.append(elem)
    table._tbl.tblPr.append(borders)


sections = [
    ("一、对大数据技术的理解", [
        "这门课学到现在，我对“大数据”的理解和刚开课时不太一样。开头我以为它就是数据量很大，最多再加上服务器多一点。后来发现这个说法太粗了。课堂里反复出现的不是“多”一个字，而是数据从哪里来、怎么存、脏数据怎么处理、最后算出来的东西能不能信。",
        "我以前写小程序，数据一般都是老师给好的，表格也比较规整，所以不会太在意前面的整理过程。学到数据清洗以后才发现，现实一点的数据经常不听话：有空值，有重复，有些字段写法还不统一。这个时候如果直接统计，结果看着像是对的，其实可能已经错了。这个印象比某个具体命令更深。",
        "HDFS 和 MapReduce 这一块，我一开始理解得比较慢。尤其是 Map 和 Reduce，刚听的时候觉得像是在绕概念。后来把它想成“先分开做小账，再把小账合起来”，就顺一点了。大数据处理不是让一台电脑从头做到尾，而是把任务拆给很多节点，这个思想我觉得挺重要。",
        "现在让我完整讲 Hadoop 生态，我肯定还讲不细，但至少知道大数据不是某一个软件名。它更像一条链子：前面数据质量差，后面分析就站不稳；中间计算设计不好，效率也会拖下来。以前我只关心代码能不能跑，现在会多想一步，数据本身有没有问题。",
    ]),
    ("二、对应用价值和就业前景的理解", [
        "大数据应用以前对我来说有点像新闻里的词。学完一些案例后再看，才发现它其实就在日常生活里。比如购物软件推荐商品，地图显示拥堵，短视频一直推类似内容，这些都不是随机发生的，背后都在收集和分析用户行为。",
        "我觉得它的价值不是简单地“让机器更聪明”，而是让很多判断有依据。商家可以看销量和用户停留时间，学校或企业可以看一些运行数据，交通系统也能根据实时情况调整。但这里面也有问题。数据多不代表一定判断准，如果指标选错了，反而会把人带偏。",
        "就业方面，我原来只知道“数据分析师”这个名字，现在才知道里面还分很多方向。有的人做数据采集和平台，有的人做清洗和仓库，有的人负责分析报告，还有人做算法模型。每个方向听起来都和数据有关，但要学的东西并不一样。",
        "如果以后真想靠近这个方向，我感觉自己还差不少。编程要更熟，数据库要补，统计知识也不能只停留在会算平均数。还有表达能力也重要。分析做完，如果说不清为什么这样算、结果说明什么，那别人也很难使用。这个问题以前我没怎么想过。",
    ]),
    ("三、个人课程收获、思考和困惑", [
        "这门课对我最大的改变，是做数据相关作业时不再那么急。以前一拿到文件就想写代码，觉得先跑出来再说。现在会先打开看一眼字段，看看有没有空值、重复值，数据类型是不是正常。这个步骤不难，但真的能少踩坑。",
        "有一次做统计，我一开始算出来的数量明显不对，但程序没有报错。我当时还以为是公式写错了，后来发现是重复记录没有处理。这个经历让我记住一件事：程序跑通和结果正确不是一回事。尤其是数据处理，错不一定会以红色报错的形式出现。",
        "我也发现自己排错能力还要练。路径、文件名、编码这些看着小，其实很容易卡住。刚开始遇到报错我会乱改，改到最后自己也不知道改了哪里。后来慢慢学会先看报错信息，再一步一步检查。这个过程有点笨，但比凭感觉改靠谱。",
        "另外，我对指标的理解也变了。以前觉得统计结果就是客观的，现在知道中间有很多选择。比如活跃用户怎么定义，是登录一次就算，还是要有实际操作？缺失值是删掉还是补上？这些选择都会影响最后结论。数据看起来冷冰冰，但处理数据的人其实一直在做判断。",
        "现在还有些问题我没想明白。课堂上的数据相对简单，真实项目里数据来源多、权限复杂，还涉及隐私保护，肯定麻烦很多。还有不同框架怎么选，什么时候用 Hadoop，什么时候用 Spark，这些我现在只能有个大概印象。",
        "写这份总结时，我也重新想了一下这门课的意义。它没有让我一下子变成会做大数据项目的人，但让我知道数据处理不是点几下工具，也不是复制一段代码。以后不管做什么方向，遇到问题先看数据、再分析原因、最后谨慎下结论，这个习惯应该是有用的。",
    ]),
]


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.54)
sec.bottom_margin = Cm(2.54)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(2.6)
doc.styles["Normal"].font.name = "宋体"
doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
doc.styles["Normal"].font.size = Pt(12)

for _ in range(5):
    p(doc)
p(doc, "《大数据处理与实践》", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
p(doc, "课程总结", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(5):
    p(doc)

table = doc.add_table(rows=6, cols=2)
table.alignment = WD_ALIGN_PARAGRAPH.CENTER
no_border(table)
for row, (label, value) in zip(table.rows, [
    ("班级", "【请填写班级】"),
    ("学号", "12303070250"),
    ("姓名", "黄彬"),
    ("教师", "成 卫"),
    ("", ""),
    ("日期", "2025年  月   日"),
]):
    row.height = Cm(0.95)
    set_cell(row.cells[0], label, bold=bool(label))
    set_cell(row.cells[1], value)
    row.cells[0].width = Cm(3.0)
    row.cells[1].width = Cm(8.5)

doc.add_page_break()
p(doc, "课程总结和心得", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
for title, paras in sections:
    p(doc, title, bold=True, before=8, after=2)
    for text in paras:
        p(doc, text, first=True)

doc.save(OUT)
plain = "\n".join(x for _, paras in sections for x in paras)
print(f"saved={OUT.resolve()}")
print(f"cn_chars={len(re.findall(r'[\u4e00-\u9fff]', plain))}")
