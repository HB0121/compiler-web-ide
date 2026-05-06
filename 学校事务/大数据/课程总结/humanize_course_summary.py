from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path("12303070250黄彬.docx")


def style_run(run, size=12, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_para(doc, text="", size=12, bold=False, align=None, first=False, before=0, after=0):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if first:
        p.paragraph_format.first_line_indent = Pt(24)
    run = p.add_run(text)
    style_run(run, size=size, bold=bold)
    return p


def set_cell(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    style_run(run, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def borderless(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement("w:" + edge)
        elem.set(qn("w:val"), "nil")
        borders.append(elem)
    table._tbl.tblPr.append(borders)


sections = [
    ("一、我对大数据技术的理解", [
        "刚开始上这门课的时候，我其实把“大数据”想得有点简单，觉得它大概就是数据很多、电脑跑得久一点。学到后面才发现，这个理解只摸到了一层皮。真正的大数据处理，麻烦的地方不只是数据量大，而是数据来源杂、格式不统一、里面还有缺失值和异常值。要想让这些数据真的能被使用，前面必须经过采集、存储、清洗、计算、分析这些环节，少一步都会影响后面的结果。",
        "我印象比较深的是分布式处理的思想。以前做普通程序时，我习惯把任务看成一台电脑上的事情；但在大数据场景里，单台机器很快就会遇到存储和计算能力的限制。HDFS、MapReduce 这些内容让我明白，很多大任务并不是靠一台机器“硬撑”，而是把数据和任务拆开，让多台机器一起完成。这个思路对我挺有启发，因为它不仅是技术方案，也像是一种解决复杂问题的方法：先拆清楚，再分头处理，最后汇总。",
        "还有一点是，我以前对数据清洗的重视不够。课程实践里，只要字段格式有一点不一致，或者路径、编码、空值没有处理好，后面的统计结果就可能变得很奇怪。有时候代码没有报错，但结果不一定就是对的。这个过程让我意识到，大数据并不是把工具打开、代码跑通就结束了，真正费时间的往往是前期检查和整理。数据质量如果不可靠，后面再漂亮的图表也只是表面好看。",
    ]),
    ("二、关于应用价值和就业前景的理解", [
        "这门课也让我感觉到，大数据离生活其实不远。平时刷购物平台、短视频、地图导航，背后都有数据在发挥作用。比如平台会根据浏览和购买记录做推荐，交通系统会根据实时路况调整路线，金融行业会通过异常交易来判断风险。这些例子以前我也听过，但学习之后再看，会更清楚它们不是凭空“智能”，而是很多数据被持续记录、整理和计算后的结果。",
        "我理解的大数据价值，首先是让判断更有依据。过去很多事情可能靠经验判断，现在可以用数据去发现趋势。比如用户为什么流失、某个商品为什么卖得好、设备什么时候可能出问题，都可以从数据里找线索。当然，数据也不是万能的。如果只看数字，不结合具体场景，也可能得出片面的结论。所以我觉得大数据的价值不只是“算得快”，还在于能帮助人把问题看得更细。",
        "从就业角度看，大数据方向的岗位比较多，但要求也没有我之前想得那么单一。数据开发更偏平台和流程，数据分析更强调业务理解和表达，算法相关岗位又需要更强的数学和编程基础。对我来说，这门课至少让我知道，如果以后想往这个方向靠，不能只停留在会用某一个软件或会写几行代码，还要把数据库、编程、统计分析和实际业务联系起来。否则做出来的结果可能只是“完成了作业”，很难真正解决问题。",
    ]),
    ("三、个人课程收获、思考和困惑", [
        "这学期学习下来，我最大的变化是处理数据时比以前更有顺序了。以前拿到数据，我经常直接开始写代码，出了问题再一处一处改。现在我会先看字段、数据类型、缺失情况，再想清楚自己到底要统计什么。这个习惯看起来很小，但对我帮助挺大，至少不会一开始就陷进报错里。",
        "课程中我也发现自己基础还不够扎实。比如有些命令或者环境配置，老师演示的时候看着不难，自己操作就容易卡住。有时候一个路径写错、一个文件名没对上，就会浪费很久。以前遇到这种情况我会有点着急，现在会先从最简单的地方查：文件在不在、字段名对不对、数据格式有没有问题。虽然排错过程不算轻松，但确实比盲目改代码有效。",
        "还有一次让我印象比较深的是，明明只是想做一个简单统计，但前面因为没有先检查重复记录，结果算出来的数量明显偏大。那时候我才发现，数据处理里很多错误不是语法错误，而是逻辑上没有想周全。程序能运行，只能说明它没有在语法上失败，不能说明分析结果就可信。",
        "我还慢慢认识到，数据分析不是只给出一个结果就行，还要能解释结果。比如同样是统计用户活跃度，按登录次数算和按实际使用时长算，得到的结论可能不一样。这个指标为什么这样定义、是否合理、有没有遗漏，这些都需要自己判断。也就是说，大数据技术看起来很偏技术，但背后其实也需要责任感。不能因为结果是计算机算出来的，就默认它一定正确。",
        "当然，我还有一些问题没有完全想明白。比如真实企业里的数据肯定比课堂样例复杂得多，数据安全和隐私到底怎么保证？如果数据分析结果和人的经验判断不一致，应该相信哪一边？不同的大数据框架在实际项目里又该怎么选？这些问题现在我还只能有一个大概认识，后面还需要通过项目和案例继续学习。",
        "总体来说，《大数据处理与实践》这门课让我不只是知道了几个概念或工具，更重要的是让我开始用数据处理的方式去看问题。以后无论是不是从事大数据相关工作，我觉得这种思维都很有用：先弄清楚问题，再整理数据，用结果支撑判断，同时也对结果保持怀疑和检查。对我来说，这就是这门课最大的收获。",
    ]),
]


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.54)
sec.bottom_margin = Cm(2.54)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(2.6)
doc.styles["Normal"].font.name = "宋体"
doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
doc.styles["Normal"].font.size = Pt(12)

for _ in range(5):
    add_para(doc)
add_para(doc, "《大数据处理与实践》", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "课程总结", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(5):
    add_para(doc)

table = doc.add_table(rows=6, cols=2)
table.alignment = WD_ALIGN_PARAGRAPH.CENTER
borderless(table)
for row, (label, value) in zip(table.rows, [
    ("班级", "【请填写班级】"),
    ("学号", "12303070250"),
    ("姓名", "黄彬"),
    ("教师", "成 卫"),
    ("", ""),
    ("日期", "2025年  月   日"),
]):
    row.height = Cm(0.95)
    set_cell(row.cells[0], label, bold=bool(label))
    set_cell(row.cells[1], value)
    row.cells[0].width = Cm(3.0)
    row.cells[1].width = Cm(8.5)

doc.add_page_break()
add_para(doc, "课程总结和心得", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

for title, paragraphs in sections:
    add_para(doc, title, bold=True, before=8, after=2)
    for text in paragraphs:
        add_para(doc, text, first=True)

doc.save(OUT)
plain = "\n".join(p for _, paras in sections for p in paras)
print(f"saved={OUT.resolve()}")
print(f"cn_chars={len(re.findall(r'[\u4e00-\u9fff]', plain))}")
