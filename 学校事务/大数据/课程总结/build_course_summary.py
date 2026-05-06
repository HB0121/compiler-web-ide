from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path("《大数据处理与实践》课程总结_可修改版.docx")


def set_cell_text(cell, text, bold=False, size=12):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_run(run, size=12, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_center_paragraph(doc, text, size=12, bold=False, space_after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    style_run(r, size=size, bold=bold)
    return p


def add_body_paragraph(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.first_line_indent = Pt(24)
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(0)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        style_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        style_run(r2)
    else:
        r = p.add_run(text)
        style_run(r)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    style_run(r, size=12, bold=True)
    return p


summary = [
    ("一、对大数据技术的理解", [
        "在学习《大数据处理与实践》之前，我对“大数据”的理解比较粗浅，更多停留在“数据量很大”“互联网公司会用”这样的印象上。经过这一阶段的学习，我慢慢意识到，大数据并不是把数据简单地堆在一起，而是一套围绕数据采集、存储、清洗、计算、分析和应用的完整技术体系。它处理的不只是数量问题，还包括数据类型复杂、产生速度快、价值密度不均衡等问题。也就是说，大数据真正难的地方不在于“多”，而在于怎样从庞杂的数据中找到可靠、有用、能够支持决策的信息。",
        "从技术层面看，我对分布式思想的印象最深。单台计算机的存储和计算能力有限，当数据规模不断扩大时，必须把任务拆分到多台机器上协同完成。Hadoop、HDFS、MapReduce 等内容让我理解了“分而治之”的基本思路：数据可以分块存储，计算任务也可以分解执行，最后再汇总结果。这个过程看似只是技术实现，背后其实体现了一种处理复杂问题的方法。面对一个庞大的任务，先拆解、再并行处理、最后整合，是大数据课程带给我的一个重要思维启发。",
        "此外，数据预处理也改变了我对数据分析的看法。以前我容易把注意力放在模型、算法或可视化结果上，认为只要工具足够强大，就能得到有价值的结论。但课程中的实践让我发现，原始数据往往存在缺失、重复、异常、格式不统一等问题，如果前期清洗不到位，后面的分析结果就可能失真。数据质量决定分析质量，这一点听起来简单，真正操作时才会发现它很考验耐心和细致程度。",
    ]),
    ("二、对大数据应用、就业前景及应用价值的理解", [
        "大数据技术的应用范围非常广。电商平台可以通过用户浏览、收藏、购买等行为进行推荐；交通管理可以根据实时路况数据优化路线和信号灯；医疗领域可以利用病例数据辅助诊断和公共卫生分析；金融行业则会用大数据进行风控、反欺诈和客户画像。通过这些案例，我感受到大数据并不是某个专业或某个行业的附属工具，而是很多行业进行数字化转型的重要基础。",
        "我认为大数据的应用价值主要体现在三个方面。第一，它能提高决策的依据性。过去很多判断依赖经验，现在可以通过数据发现趋势和规律。第二，它能提升服务的精准度，例如个性化推荐、智能客服、精准营销等，都是在理解用户行为的基础上实现的。第三，它能帮助发现潜在风险，比如异常交易识别、设备故障预测、舆情监测等，这些都体现了数据在预警和管理中的作用。",
        "从就业前景来看，大数据相关岗位既有技术型，也有分析型。数据开发工程师更关注数据平台、数据仓库和计算流程；数据分析师更强调业务理解、指标设计和结果解释；算法工程师则偏向模型构建和优化。对我们学生来说，这门课让我看到一个现实问题：只会使用工具还不够，未来如果想在相关方向发展，还需要同时具备编程能力、统计思维、业务理解和表达能力。尤其是数据分析结果最终要服务于实际问题，不能只停留在代码运行成功或者图表画出来。",
    ]),
    ("三、个人课程收获与反思", [
        "这门课给我的最大收获，是让我对数据处理流程有了更完整的认识。以前做一些小作业时，我常常拿到数据就直接开始写代码，遇到问题再临时修改。现在我会先观察字段含义、数据类型和缺失情况，再考虑清洗规则和分析目标。这个变化虽然不算特别显眼，但对我来说很重要，因为它让我从“为了完成任务而操作”逐渐转向“为了回答问题而处理数据”。",
        "在实践过程中，我也发现自己有不少不足。比如对一些命令和工具的使用还不够熟练，遇到报错时容易先怀疑环境配置，而不是冷静地从数据格式、路径、依赖和逻辑顺序逐步排查。有时候代码能跑通，但我对每一步为什么这样做解释得不够清楚。课程让我意识到，真正掌握一项技术，不只是会按照步骤操作，还要能说清楚它适合解决什么问题、有什么限制、结果是否可信。",
        "我还有一个比较深的感受是，大数据技术虽然强调效率和规模，但最后仍然离不开人的判断。数据本身不会自动告诉我们答案，指标怎么定义、异常值怎么处理、结果如何解释，都需要结合场景。比如同样是用户活跃度，不同业务可能有完全不同的衡量方式；同样是推荐系统，如果只追求点击率，也可能忽视用户体验或信息茧房问题。因此，学习大数据技术时不能只关注“能不能做”，还要思考“该不该这样做”和“这样做会带来什么影响”。",
        "当然，这门课也留下了一些困惑。比如在真实企业场景中，数据规模更大、数据来源更多，如何保证数据安全和隐私保护？当数据分析结果和人的经验判断发生冲突时，应该如何权衡？不同大数据平台和计算框架之间应该怎样选择？这些问题我现在还不能完全回答，但它们让我意识到后续学习不能只停留在课本内容，还需要通过项目实践和案例阅读继续深化。",
        "总体来说，《大数据处理与实践》让我从概念、工具和实践三个层面重新认识了数据。它不仅让我了解了大数据技术的基本框架，也让我意识到数据处理是一项需要逻辑、耐心和责任感的工作。今后无论是否直接从事大数据岗位，我都希望能保持数据意识：面对问题时尽量用数据支撑判断，面对数据时也保持审慎，不盲目相信表面结果。对我来说，这门课的意义不只是完成一次课程学习，而是帮助我建立起一种更清晰、更理性的分析习惯。",
    ]),
]


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.54)
section.bottom_margin = Cm(2.54)
section.left_margin = Cm(3.0)
section.right_margin = Cm(2.6)

styles = doc.styles
styles["Normal"].font.name = "宋体"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Normal"].font.size = Pt(12)

for _ in range(4):
    doc.add_paragraph()
add_center_paragraph(doc, "《大数据处理与实践》", size=18, bold=True, space_after=12)
add_center_paragraph(doc, "课程总结与心得", size=22, bold=True, space_after=48)
for _ in range(3):
    doc.add_paragraph()

table = doc.add_table(rows=4, cols=2)
table.alignment = WD_ALIGN_PARAGRAPH.CENTER
table.style = "Table Grid"
labels = ["课程名称", "班级", "学号", "姓名"]
values = ["大数据处理与实践", "【请填写班级】", "【请填写学号】", "【请填写姓名】"]
for row, label, value in zip(table.rows, labels, values):
    row.height = Cm(1.1)
    set_cell_text(row.cells[0], label, bold=True)
    set_cell_text(row.cells[1], value)
    set_cell_shading(row.cells[0], "F2F2F2")
    row.cells[0].width = Cm(4.0)
    row.cells[1].width = Cm(8.5)

for _ in range(8):
    doc.add_paragraph()
add_center_paragraph(doc, "说明：提交前请将班级、学号、姓名替换为本人信息，并结合自己的课堂实践补充一两处真实经历。", size=10)

doc.add_page_break()
add_center_paragraph(doc, "课程总结和心得", size=16, bold=True, space_after=12)

for heading, paras in summary:
    add_heading(doc, heading)
    for para in paras:
        add_body_paragraph(doc, para)

doc.add_page_break()
add_center_paragraph(doc, "原作业要求", size=14, bold=True, space_after=8)
requirements = [
    "1、对大数据技术的理解",
    "2、对大数据应用和就业前景及应用价值的理解",
    "3、个人课程收获（总结的重点：所得、所思、所惑。。。）",
    "4、总结至少1500字，内容必须完整",
    "5、文档格式要求：宋体，小四号字，1.5倍行距，首行缩进2个字",
    "6、封面有课程名称、班级、学号、姓名；A4纸单面打印",
]
for req in requirements:
    add_body_paragraph(doc, req)

doc.save(OUT)

plain = "\n".join(p for _, paras in summary for p in paras)
cn_chars = len(re.findall(r"[\u4e00-\u9fff]", plain))
print(f"saved={OUT.resolve()}")
print(f"chinese_chars={cn_chars}")
