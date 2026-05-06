from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


path = "《大数据处理与实践》课程总结_提交版.docx"
doc = Document(path)


def style_run(run, size=12, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def clear_para(paragraph):
    for run in paragraph.runs:
        run.text = ""


def set_para(paragraph, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    clear_para(paragraph)
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    style_run(run, size=size, bold=bold)


def set_cell(cell, text, size=12, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    style_run(run, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_borderless(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


for paragraph in doc.paragraphs[:12]:
    clear_para(paragraph)

set_para(doc.paragraphs[4], "《大数据处理与实践》", size=22, bold=True)
set_para(doc.paragraphs[5], "课程总结", size=22, bold=True)

table = doc.tables[0]
while len(table.rows) < 6:
    table.add_row()

items = [
    ("班级", "【请填写班级】"),
    ("学号", "【请填写学号】"),
    ("姓名", "【请填写姓名】"),
    ("教师", "成 卫"),
    ("", ""),
    ("日期", "2025年  月   日"),
]

for row, (label, value) in zip(table.rows, items):
    row.height = Cm(0.95)
    cells = row.cells
    set_cell(cells[0], label, bold=bool(label))
    set_cell(cells[1], value)
    cells[0].width = Cm(3.0)
    cells[1].width = Cm(8.5)

set_borderless(table)
doc.save(path)
print(path)
