from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = "实验三_健身房会员管理系统_软件项目解决方案及项目计划.docx"


def set_run_font(run, name="宋体", size=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph(paragraph, align=None, first_line=None, space_after=6, line_spacing=1.25):
    if align is not None:
        paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing
    if first_line is not None:
        fmt.first_line_indent = Cm(first_line)


def add_text(paragraph, text, size=12, bold=False, font="宋体"):
    run = paragraph.add_run(text)
    set_run_font(run, font, size, bold)
    return run


def add_body(doc, text):
    p = doc.add_paragraph()
    set_paragraph(p, first_line=0.74, space_after=6, line_spacing=1.3)
    add_text(p, text, 12)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading("", level=level)
    set_paragraph(p, space_after=8 if level == 1 else 6, line_spacing=1.2)
    run = p.add_run(text)
    set_run_font(run, "黑体", 16 if level == 1 else 14, True)
    run.font.color.rgb = RGBColor(31, 78, 121) if level == 1 else RGBColor(68, 68, 68)
    return p


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, "宋体", 10.5, bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_text(hdr[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(hdr[i], "D9EAF7")
        if widths:
            hdr[i].width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, item in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if len(str(item)) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], str(item), align=align)
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style=None)
        set_paragraph(p, first_line=0, space_after=4, line_spacing=1.25)
        add_text(p, "（{}）".format(item[0]), 12, bold=True)
        add_text(p, item[1], 12)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.0)
    sec.right_margin = Cm(2.6)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(12)

    # Cover
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "重庆理工大学", 20, True, "黑体")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "计算机科学与工程学院", 18, True, "黑体")
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "软件工程实验报告", 24, True, "黑体")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "实验三 软件项目解决方案及项目计划", 18, True, "黑体")
    for _ in range(4):
        doc.add_paragraph()
    info = [
        ("项目名称", "健身房会员管理系统"),
        ("文档名称", "软件项目解决方案及项目计划"),
        ("班级", "123030701"),
        ("学号", "12303070250"),
        ("姓名", "黄转"),
        ("日期", "2026年5月8日"),
    ]
    table = doc.add_table(rows=len(info), cols=2)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row, (k, v) in zip(table.rows, info):
        row.cells[0].width = Cm(4)
        row.cells[1].width = Cm(8)
        set_cell_text(row.cells[0], k, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(row.cells[1], v, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    # Static TOC
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "目录", 18, True, "黑体")
    toc_items = [
        "1 项目概述",
        "2 系统解决方案设计",
        "3 核心问题解决方案",
        "4 项目开发工作量估算",
        "5 项目开发计划",
        "6 项目风险与保障措施",
        "7 结论",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        set_paragraph(p, first_line=0.5, space_after=4, line_spacing=1.2)
        add_text(p, item, 12)
    doc.add_page_break()

    add_heading(doc, "1 项目概述")
    add_body(doc, "健身房会员管理系统是面向健身房日常运营场景的信息管理系统，主要服务对象包括前台工作人员、门店管理人员和系统管理员。系统拟对会员建档、办卡缴费、到店签到、续费提醒、经营统计等业务进行统一管理，用信息化方式替代纸质登记、分散表格和人工统计。")
    add_body(doc, "根据前期问题定义与可行性研究，当前业务中存在会员信息分散、签到记录不完整、到期提醒滞后、经营数据统计不及时、人工操作容易出错等问题。实验三在实验二分析结果基础上，进一步给出系统硬件、网络、信息安全、核心业务处理方案，并进行开发工作量估算与五个月项目计划设计。")

    add_table(
        doc,
        ["用户角色", "主要职责", "典型操作"],
        [
            ["前台工作人员", "办理会员业务并维护基础数据", "新会员登记、办卡缴费、签到确认、续费提醒"],
            ["门店管理人员", "掌握经营情况并进行业务决策", "查看会员规模、收入统计、活跃度、续费情况"],
            ["系统管理员", "维护系统运行环境和权限配置", "账号管理、权限分配、数据备份、日志检查"],
        ],
        [3.2, 5.6, 6.4],
    )

    add_heading(doc, "2 系统解决方案设计")
    add_heading(doc, "2.1 硬件解决方案", 2)
    add_body(doc, "系统采用小型门店适用的 B/S 架构部署方式。前台通过普通办公电脑或平板浏览器访问系统，管理端通过内网或安全公网入口访问。后端可部署在云服务器，也可部署在门店本地服务器；考虑课程实验和中小型健身房成本，推荐采用云服务器加门店终端的轻量方案。")
    add_table(
        doc,
        ["硬件类别", "建议配置", "用途说明"],
        [
            ["应用/数据库服务器", "2 核 CPU、4GB 内存、80GB SSD，可扩展云盘", "部署 Web 服务、数据库、备份任务，支撑 1-3 家门店初期运行"],
            ["前台终端", "Windows 办公电脑或平板，浏览器访问", "会员登记、办卡缴费、签到确认和查询"],
            ["签到设备", "二维码扫码器或会员卡读卡器，可保留手机号手工签到", "提升签到效率，设备异常时仍可人工录入"],
            ["网络设备", "千兆路由器、交换机、稳定宽带", "保障前台终端与云端服务稳定连接"],
            ["备份设备", "云备份空间或移动硬盘", "保存数据库备份、导出报表和系统日志"],
        ],
        [3.4, 5.6, 6.2],
    )

    add_heading(doc, "2.2 软件与部署方案", 2)
    add_body(doc, "系统软件建议采用分层结构：表现层负责页面交互，业务层处理会员、缴费、签到、提醒和统计逻辑，数据层负责会员档案、会员卡、缴费记录、签到记录、用户权限和日志等数据的持久化。数据库采用关系型数据库，便于保证交易记录一致性和后续统计查询。")
    add_table(
        doc,
        ["层次", "主要内容", "设计要点"],
        [
            ["表现层", "前台业务页面、管理统计页面、系统设置页面", "界面简洁，常用操作入口明显，支持按手机号/卡号快速检索"],
            ["业务层", "会员建档、办卡缴费、签到、续费提醒、统计分析", "业务规则集中处理，避免重复录入和状态不一致"],
            ["数据层", "会员表、会员卡表、缴费表、签到表、用户表、日志表", "设置主键、唯一约束和必要索引，保证查询效率和数据准确性"],
            ["支撑层", "备份、权限、日志、异常告警", "保障系统稳定运行，便于问题追踪与恢复"],
        ],
        [2.6, 5.8, 6.8],
    )

    add_heading(doc, "2.3 网络设计", 2)
    add_body(doc, "系统网络结构以稳定、低成本和易维护为原则。门店内前台终端、扫码器和管理电脑接入同一局域网，再通过宽带访问云端服务。云端服务器开放 HTTPS 服务端口，数据库不直接暴露到公网。若采用本地服务器，则前台终端通过内网访问，本地服务器定时向云端或外部存储同步备份。")
    add_table(
        doc,
        ["网络环节", "设计方案", "说明"],
        [
            ["门店局域网", "前台电脑、管理电脑、扫码设备接入路由器/交换机", "保证签到和办卡操作在门店内快速完成"],
            ["外网访问", "通过 HTTPS 访问云端系统", "防止账号、会员资料和缴费数据明文传输"],
            ["数据库访问", "仅允许应用服务器内网访问数据库", "降低数据库被直接攻击的风险"],
            ["备份链路", "每日自动备份到云存储或异地空间", "防止设备损坏、误操作导致数据丢失"],
        ],
        [3.2, 6.2, 5.8],
    )

    add_heading(doc, "2.4 信息安全方案", 2)
    add_bullets(
        doc,
        [
            ("身份认证", "系统登录采用账号密码认证，关键岗位账号由管理员统一创建，离职人员账号及时停用。"),
            ("权限控制", "按前台、管理人员、系统管理员设置角色权限，前台只能办理业务和查询必要信息，管理人员可查看统计报表，管理员负责系统配置。"),
            ("数据保护", "会员手机号、身份证号等敏感字段限制展示范围，数据库定期备份，重要操作写入日志。"),
            ("传输安全", "系统部署启用 HTTPS，避免会员资料、缴费数据和登录凭据在传输过程中泄露。"),
            ("应急恢复", "制定数据备份与恢复流程，发生误删、服务器故障或网络中断时可按最近备份恢复。"),
        ],
    )

    add_heading(doc, "3 核心问题解决方案")
    add_body(doc, "本系统的核心问题不是单纯录入数据，而是要让会员生命周期、缴费有效期、签到记录和经营统计保持一致。针对实验二提出的业务痛点，给出如下解决方案。")
    add_table(
        doc,
        ["核心问题", "解决方案", "预期效果"],
        [
            ["会员信息分散", "建立统一会员档案，以手机号或会员号作为唯一检索入口，会员卡、缴费、签到记录关联会员档案", "减少重复录入，保证基础信息一致"],
            ["办卡缴费效率低", "设计标准办卡流程：录入资料、选择会员类型、生成卡号、登记缴费、更新有效期", "前台可按固定步骤快速办理，降低漏项风险"],
            ["签到统计困难", "支持卡号、手机号或二维码签到，签到时自动校验会员有效期并记录时间", "可快速统计到店次数、活跃会员和异常签到"],
            ["续费提醒滞后", "系统每日扫描有效期，提前 7 天、3 天和到期当天生成提醒清单", "提高续费跟进及时性，减少会员流失"],
            ["经营分析不及时", "按日、周、月统计新增会员、续费金额、到店次数和活跃率", "为课程安排、促销活动和人员排班提供数据依据"],
            ["断网或设备异常", "保留手机号手工查询与离线登记表，网络恢复后由前台补录并标记来源", "保证门店业务不中断，避免因设备问题影响服务"],
        ],
        [3.2, 7.0, 5.0],
    )

    add_heading(doc, "4 项目开发工作量估算")
    add_body(doc, "根据系统规模、功能模块和课程实验要求，本项目采用按阶段分解的工作量估算法。估算单位为人月，综合考虑需求分析、概要设计、详细设计、编码、测试、部署和项目管理等活动。系统属于中小型信息管理系统，功能复杂度中等，数据处理以增删改查、状态更新和统计查询为主。")
    add_table(
        doc,
        ["工作阶段", "主要工作内容", "工作量（人月）", "占比"],
        [
            ["需求分析", "梳理业务流程、用户角色、数据项、功能边界，形成需求说明", "1.2", "12%"],
            ["概要设计", "确定总体架构、模块划分、数据库概念结构和安全方案", "1.0", "10%"],
            ["详细设计", "设计数据库表、接口、页面流程、权限规则和异常处理", "1.2", "12%"],
            ["编码实现", "实现会员、缴费、签到、提醒、统计、权限和日志等模块", "3.5", "35%"],
            ["系统测试", "完成单元测试、集成测试、业务流程测试和缺陷修复", "1.4", "14%"],
            ["部署集成", "完成软件部署、数据库初始化、硬件/扫码设备接入和数据导入", "0.9", "9%"],
            ["培训与项目管理", "用户培训、文档编写、进度协调、验收准备", "0.8", "8%"],
        ],
        [3.0, 6.4, 2.6, 2.0],
    )
    add_body(doc, "由上表可得项目总工作量约为 10.0 人月。若项目计划周期为 5 个月，则平均需要 2 名全职人员持续投入。考虑不同阶段工作并行和角色差异，实际团队可由项目经理/需求分析人员、后端开发、前端开发、测试与实施人员组成，小组成员可兼职承担部分角色。")

    add_table(
        doc,
        ["角色", "投入人月", "主要职责"],
        [
            ["项目经理/需求分析", "1.4", "需求访谈、范围控制、计划协调、验收沟通"],
            ["架构与后端开发", "2.6", "系统架构、数据库设计、核心业务接口和安全控制"],
            ["前端开发", "2.0", "前台业务页面、管理报表页面、交互优化"],
            ["测试工程师", "1.4", "测试用例、集成测试、回归测试和缺陷跟踪"],
            ["实施与运维", "1.0", "服务器部署、硬件接入、数据迁移、备份配置"],
            ["文档与培训", "0.8", "用户手册、培训材料、验收文档"],
            ["项目管理预留", "0.8", "会议、评审、风险处理和进度缓冲"],
        ],
        [3.2, 2.8, 8.0],
    )

    add_heading(doc, "5 项目开发计划")
    add_body(doc, "用户要求项目在 5 个月内完成，因此采用迭代式计划安排。前 2 个月完成需求和设计，第 3 至第 4 个月完成主体开发与集成测试，第 5 个月完成部署、培训、试运行和验收。")
    add_table(
        doc,
        ["月份", "阶段目标", "主要任务", "阶段成果"],
        [
            ["第1个月", "需求分析与计划启动", "确认业务范围、用户角色、功能清单、数据项和项目计划；完成原型草图", "需求规格说明、项目计划、初步原型"],
            ["第2个月", "总体设计与详细设计", "完成系统架构、数据库设计、页面流程、权限设计、硬件和网络方案", "概要设计说明、数据库设计、接口清单"],
            ["第3个月", "核心功能开发", "实现会员档案、办卡缴费、签到、权限登录和基础查询功能", "可运行的核心业务版本"],
            ["第4个月", "扩展功能与系统测试", "实现续费提醒、统计报表、日志备份；开展集成测试和缺陷修复", "测试报告、修复后的候选版本"],
            ["第5个月", "部署集成与验收", "完成服务器部署、数据初始化、扫码设备接入、用户培训、试运行和验收", "上线版本、用户手册、验收材料"],
        ],
        [2.2, 3.4, 7.0, 4.2],
    )

    add_table(
        doc,
        ["任务", "第1月", "第2月", "第3月", "第4月", "第5月"],
        [
            ["需求分析", "●", "", "", "", ""],
            ["概要/详细设计", "●", "●", "", "", ""],
            ["数据库与后端开发", "", "●", "●", "●", ""],
            ["前端页面开发", "", "●", "●", "●", ""],
            ["硬件与软件集成", "", "", "", "●", "●"],
            ["系统测试与缺陷修复", "", "", "●", "●", "●"],
            ["培训、试运行、验收", "", "", "", "", "●"],
        ],
        [3.2, 2.2, 2.2, 2.2, 2.2, 2.2],
    )

    add_heading(doc, "6 项目风险与保障措施")
    add_table(
        doc,
        ["风险类别", "风险描述", "保障措施"],
        [
            ["需求风险", "会员类型、收费规则或续费规则变化导致返工", "需求阶段确认规则清单，重要变更走评审流程"],
            ["数据风险", "历史会员资料不完整或导入错误", "上线前整理模板，导入后抽样核对，保留原始数据备份"],
            ["安全风险", "会员隐私信息泄露或越权访问", "角色权限、日志审计、HTTPS、数据库备份和敏感信息限制展示"],
            ["进度风险", "开发任务集中在后期导致测试不足", "按月交付阶段成果，核心功能优先开发，预留测试修复时间"],
            ["设备风险", "扫码器、网络或服务器故障影响签到", "保留手工签到和补录机制，部署监控与定期备份"],
        ],
        [3.0, 5.8, 6.4],
    )

    add_heading(doc, "7 结论")
    add_body(doc, "健身房会员管理系统具有明确的业务需求和较好的技术、经济、操作可行性。通过统一会员档案、标准化办卡缴费流程、自动签到记录、续费提醒和经营统计，可以有效解决当前人工管理效率低、数据分散、提醒不及时和统计困难等问题。")
    add_body(doc, "在 5 个月开发周期内，项目总工作量约 10.0 人月，采用 2 人左右的平均投入并按阶段安排需求、设计、开发、测试、部署和验收工作，能够满足课程实验对软件项目解决方案和项目计划设计的要求。")

    doc.save(OUT)


if __name__ == "__main__":
    build()
